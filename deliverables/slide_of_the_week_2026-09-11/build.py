"""Slide of the week, 11 Sep 2026 (Ann 1:1): figure + pptx, house style, words over numbers.

One slide, one 2 x 3 figure, one column per Addis season:
  top    = Addis FTIR EC from that season's own calibration vs HIPS Fabs / MAC 10, Deming line,
           1:1, a few words saying what went right and wrong (numbers live in the notes);
  bottom = every spectrum that calibration trained on (grey) against every Addis spectrum
           in that season (colour), same axes: the training set is far more loaded than Addis.

Run with anaconda python from the repo root:
    ~/anaconda3/bin/python deliverables/slide_of_the_week_2026-09-11/build.py

Inputs (the 10 Sep weekly run that was shown in the meeting; nothing refitted):
- research/ftir_hips_chem/output/tables/ann_weekly_20260910/{predictions,regression_metrics,
  addis_evaluation,seasonal_overlap,cohort_no_co2__<season>}.csv
- research/ftir_ec_phase3/output/corrected/{improve_pool,etad}_corrected_df6.npz (AIRSpec df 6)
"""

from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
FIG = HERE / "figures"
FIG.mkdir(exist_ok=True)
T = ROOT / "research/ftir_hips_chem/output/tables/ann_weekly_20260910"
CORR = ROOT / "research/ftir_ec_phase3/output/corrected"

GREY, BLUE, PURPLE, ACCENT, AMBER, INK = (
    "#8F8C84", "#2C6E9E", "#7A4FA3", "#B23327", "#C49442", "#22252A")
DPI = 168
plt.rcParams.update({
    "figure.facecolor": "white", "axes.facecolor": "white",
    "axes.spines.top": False, "axes.spines.right": False,
    "axes.edgecolor": INK, "axes.labelcolor": INK, "xtick.color": INK,
    "ytick.color": INK, "text.color": INK, "font.size": 11,
    "axes.titlesize": 12.5, "legend.frameon": False,
})
MINUS = "−"

SEASONS = [
    ("Dry (Oct-Feb)", "Dry__Oct_Feb", "Dry season", AMBER,
     "Offset gone,\nbut the slope collapses"),
    ("Belg (Mar-May)", "Belg__Mar_May", "Belg (short rains)", PURPLE,
     "Shallow and scattered,\nneither problem solved"),
    ("Kiremt (Jun-Sep)", "Kiremt__Jun_Sep", "Kiremt (long rains)", BLUE,
     "Slope close to 1:1,\nbut it all reads low"),
]


def save(fig, name):
    p = FIG / name
    fig.savefig(p, dpi=DPI, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    Image.open(p).convert("RGB").save(p)
    return p


def deming(x, y, lam):
    mx, my = x.mean(), y.mean()
    sxx, syy = ((x - mx) ** 2).mean(), ((y - my) ** 2).mean()
    sxy = ((x - mx) * (y - my)).mean()
    d = syy - lam * sxx
    m = (d + np.sqrt(d * d + 4 * lam * sxy * sxy)) / (2 * sxy)
    return m, my - m * mx


def ci(v, lo, hi, nd=2):
    f = lambda z: f"{z:.{nd}f}".replace("-", MINUS)
    return f"{f(v)} [{f(lo)}, {f(hi)}]"


# ---------------------------------------------------------------- data
pred = pd.read_csv(T / "predictions.csv")
met = pd.read_csv(T / "regression_metrics.csv")
ev = pd.read_csv(T / "addis_evaluation.csv")
overlap = pd.read_csv(T / "seasonal_overlap.csv")
lib = np.load(CORR / "improve_pool_corrected_df6.npz", allow_pickle=True)
tgt = np.load(CORR / "etad_corrected_df6.npz", allow_pickle=True)
wn = lib["wn"].astype(float)
pos = pd.Series(np.arange(len(lib["analysis_id"])), index=lib["analysis_id"])
tmid = tgt["media_id"].astype(int)
X_addis = np.vstack([tgt["corrected"][tmid == m].mean(axis=0) for m in ev.MediaId])  # replicate scans of one filter

rows = {}
for key, slug, *_ in SEASONS:
    fid = f"no_co2__{slug}"
    p = pred[(pred.fit_id == fid) & (pred.season == key)]
    r = met[(met.fit_id == fid) & (met.evaluation_group == key)].iloc[0]
    m, b = deming(p.hips_equivalent_ugm3.values, p.prediction_ugm3.values, r.deming_lambda)
    # the meeting's numbers must reproduce from the saved predictions before anything is drawn
    assert np.allclose([m, b], [r.deming_slope, r.deming_intercept], atol=1e-6), (fid, m, b)
    cohort = pd.read_csv(T / f"cohort_no_co2__{slug}.csv")
    train = lib["corrected"][pos.loc[cohort.loc[cohort.role.eq("train"), "AnalysisId"]].to_numpy()]
    rows[key] = dict(p=p, r=r, m=m, b=b, train=train, addis=X_addis[ev.season.eq(key).to_numpy()])
assert [len(rows[k]["train"]) for k, *_ in SEASONS] == [394, 373, 390]
hist = met[(met.fit_id == "historical_ocec440__All_Addis") & (met.evaluation_group == "All Addis")].iloc[0]

# ---------------------------------------------------------------- figure
fig, ax = plt.subplots(2, 3, figsize=(14.2, 7.4),
                       gridspec_kw={"height_ratios": [1.15, 1], "hspace": 0.42, "wspace": 0.22})
hi = float(np.ceil(max(max(d["p"].hips_equivalent_ugm3.max(), d["p"].prediction_ugm3.max())
                       for d in rows.values()) + 0.5))
ymax_spec = max(np.nanmax(d["train"]) for d in rows.values()) * 1.04
co2 = (1800, 2500)

for j, (key, slug, name, col, verdict) in enumerate(SEASONS):
    d = rows[key]
    a = ax[0, j]
    xx = np.array([0, hi])
    a.plot(xx, xx, ls=":", color=GREY, lw=1.3, zorder=1)
    a.scatter(d["p"].hips_equivalent_ugm3, d["p"].prediction_ugm3, s=20, color=col, alpha=0.75, lw=0, zorder=3)
    a.plot(xx, d["m"] * xx + d["b"], color=INK, lw=1.8, zorder=4)
    a.set_xlim(0, hi)
    a.set_ylim(0, hi)
    a.set_aspect("equal")
    a.set_title(name, loc="left", color=col, fontweight="bold")
    a.text(0.04, 0.96, verdict, transform=a.transAxes, va="top", fontsize=11.5, color=INK)
    a.set_xlabel("HIPS Fabs / MAC 10 (µg/m³)")
    if j == 0:
        a.set_ylabel("FTIR EC, own-season\ncalibration (µg/m³)")
        a.plot([], [], ls=":", color=GREY, lw=1.3, label="1:1")
        a.plot([], [], color=INK, lw=1.8, label="Deming fit")
        a.legend(loc="center right", fontsize=9.5)

    s = ax[1, j]
    s.axvspan(*co2, color="#EEECE6", lw=0, zorder=0)
    s.plot(wn, d["train"].T, color=GREY, alpha=0.12, lw=0.6, zorder=1)
    s.plot(wn, d["addis"].T, color=col, alpha=0.55, lw=0.7, zorder=2)
    s.set_xlim(4000, 1426)
    s.set_ylim(-0.004, ymax_spec)
    s.set_yticks([])
    s.spines["left"].set_visible(False)
    s.set_xlabel("Wavenumber (cm⁻¹)")
    if j == 0:
        s.plot([], [], color=GREY, lw=2, label="IMPROVE filters it trained on")
        s.plot([], [], color=AMBER, lw=2, label="Addis filters, this season")
        s.legend(loc="upper right", fontsize=9.5, handlelength=1.2)
        s.set_ylabel("Absorbance (spline baseline)")
        s.text(2150, ymax_spec * 0.55, "not used\nfor matching", ha="center", fontsize=8.8, color=GREY)
    else:
        s.text(0.97, 0.95, "Addis sits at the bottom\nof the training range", transform=s.transAxes,
               ha="right", va="top", fontsize=10, color=INK)

fig.text(0.075, 0.455, "Likely why: the analogs match Addis in shape, but span a much wider loading range",
         fontsize=12.5, fontweight="bold", color=INK)
fig_path = save(fig, "f1_seasonal_calibrations_and_training_spectra.png")

# ---------------------------------------------------------------- pptx
INK_RGB = RGBColor(0x22, 0x25, 0x2A)
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
sl = prs.slides.add_slide(prs.slide_layouts[6])
tb = sl.shapes.add_textbox(Inches(0.45), Inches(0.22), Inches(12.5), Inches(0.95))
tb.text_frame.word_wrap = True
p = tb.text_frame.paragraphs[0]
p.text = ("A calibration per season fixes the offset or the slope, never both, "
          "and the analogs span a much wider loading range than Addis")
p.font.size, p.font.bold, p.font.color.rgb = Pt(20), True, INK_RGB
im = Image.open(fig_path)
scale = min(12.4 / (im.width / DPI), 6.05 / (im.height / DPI), 1.35)
w = im.width / DPI * scale
sl.shapes.add_picture(str(fig_path), Inches((13.333 - w) / 2),
                      Inches(1.2 + (6.05 - im.height / DPI * scale) / 2), Inches(w))

o = overlap[overlap["mask"] == "no_co2"].set_index(["group_a", "group_b"]).intersection
ov = (int(o["Dry (Oct-Feb)", "Belg (Mar-May)"]), int(o["Dry (Oct-Feb)", "Kiremt (Jun-Sep)"]),
      int(o["Belg (Mar-May)", "Kiremt (Jun-Sep)"]))
dr, be, ki = (rows[k]["r"] for k, *_ in SEASONS)
say = (
    "This week we tried giving each Addis season its own calibration, picking the IMPROVE filters whose "
    "spectra look most like that season. Each column is one season. Top row: what that calibration gives "
    "for Addis against HIPS, with the dotted one-to-one line. In the dry season the offset disappears, the line "
    "goes through zero, but it is far too flat, so FTIR badly underestimates. Belg is in between and noisy. "
    "Kiremt, the long rains with the most charcoal burning, gets the slope about right but the whole line still "
    "sits low. So we can move the problem around, but no season fixes both at once. The bottom row is the "
    "likely reason. Grey is every IMPROVE spectrum each calibration trained on, colour is the Addis spectra for "
    "that season, on the same scale. Addis sits in a narrow band near the bottom; the training sets reach much "
    "higher, most clearly in the two rainy seasons. We matched on shape, not on how loaded the filter is, so the model is tuned to a much wider range "
    "than it is asked to measure. Next step: keep the shape matching, then keep only the analogs whose "
    "absorbance looks like Addis, cutting where the spread jumps."
)
notes = (
    "Source: the 10 Sep weekly run shown in the meeting (research/ftir_hips_chem/output/tables/ann_weekly_20260910, "
    "builder workflows/run_ann_weekly_20260910.py); nothing refitted here, and this script asserts the Deming fits "
    "reproduce from predictions.csv. Setup: analog selection by correlation with the season's median Addis "
    "spectrum on spline-baselined (AIRSpec df 6) spectra with the CO2 band 1800 to 2500 cm-1 removed from matching "
    "only; top 500 TOR-eligible IMPROVE filters; PLS on the full 1426 to 3998 cm-1 grid; Option A style split "
    "(20% of sites held out for the TOR test, site-grouped 5-fold CV, first major minimum). MAC 10. "
    f"Deming λ {dr.deming_lambda:.2f} (HIPS uncertainty vs historical TOR RMSE), 95% month-block bootstrap. "
    f"Dry: slope {ci(dr.deming_slope, dr.slope_ci_low, dr.slope_ci_high)}, intercept "
    f"{ci(dr.deming_intercept, dr.intercept_ci_low, dr.intercept_ci_high)}, n {int(dr.n)}, k {int(dr.k)}. "
    f"Belg: slope {ci(be.deming_slope, be.slope_ci_low, be.slope_ci_high)}, intercept "
    f"{ci(be.deming_intercept, be.intercept_ci_low, be.intercept_ci_high)}, n {int(be.n)}, k {int(be.k)}. "
    f"Kiremt: slope {ci(ki.deming_slope, ki.slope_ci_low, ki.slope_ci_high)}, intercept "
    f"{ci(ki.deming_intercept, ki.intercept_ci_low, ki.intercept_ci_high)}, n {int(ki.n)}, k {int(ki.k)}. "
    f"Test set TOR R²: Dry {dr.TOR_R2:.2f}, Belg {be.TOR_R2:.2f}, Kiremt {ki.TOR_R2:.2f}, all below the 0.85 "
    f"screen, so none is adopted; reference stays the Lowest-OC/EC 440 calibration (k 8, Addis "
    f"{hist.deming_slope:.2f}x {MINUS} {abs(hist.deming_intercept):.2f}, TOR R² {hist.TOR_R2:.2f}). "
    f"Analog overlap between seasons (of 500): Dry and Belg share {ov[0]}, Dry and Kiremt {ov[1]}, Belg and Kiremt {ov[2]}, which "
    "fits the PMF picture of a different dominant source in the dry season. Removing CO2 from matching swaps 22 "
    "of 500 analogs; also dropping above 3600 or 3500 cm-1 swaps 61 or 70. Spectra are every individual filter, "
    "not averages; replicate Addis scans of one filter are averaged. Training sets: 394 / 373 / 390 filters. "
    "Raising MAC would steepen the dry slope but push Addis concentrations implausibly low, so it is not the fix. "
    "Ann's asks from this meeting: report slope and intercept as value [low, high]; also drop >3500 cm-1 from the "
    "PLS fit itself, not just from matching; after correlation, keep the 50 to 100 analogs closest in absorbance "
    "(plot variance against rank and cut before the jump); try PMF with wood + charcoal and sea salt + polluted "
    "marine combined; move the spectra to the start of the talk; seasons first for Satoshi next Thursday. "
    "Adama (separate thread): 5 quartz TOR filters, July 2024, OC/EC mid-IMPROVE and FTIR agrees with TOR, so "
    "Adama and Bishoftu don't stand in for Addis; Ann will ask JPL for quartz sampling in Addis and pull "
    "Christian's co-located Delhi quartz/FTIR. To do: Adama HIPS on the same plot (two y axes), FTIR vs Fabs "
    "coloured by sampled volume, PurpleAir or other PM for Adama/Bishoftu vs Addis."
)
sl.notes_slide.notes_text_frame.text = f"SAY: {say}\n\nNOTES: {notes}"

out = HERE / "slide_of_the_week_2026-09-11.pptx"
prs.save(out)

# ---------------------------------------------------------------- checks
chk = Presentation(out)
assert len(chk.slides) == 1
for s_ in chk.slides:
    txt = " ".join(sh.text_frame.text for sh in s_.shapes if sh.has_text_frame)
    nt = s_.notes_slide.notes_text_frame.text
    assert chr(0x2014) not in txt + nt, "em dash on slide"
    assert nt.startswith("SAY:") and "NOTES:" in nt, "notes missing"
    assert sum(1 for sh in s_.shapes if sh.shape_type == 13) == 1, "figure missing"
assert chr(0x2014) not in Path(__file__).read_text(), "em dash in builder"
print(f"saved {out.relative_to(ROOT)}: 1 slide, 1 figure, notes present, no em dashes")
