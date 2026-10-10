"""Slide of the week, 25 Sep 2026 (Ann 1:1, gallery walkthrough): figure + pptx, house style.

One slide, one two-panel figure:
  left  = Addis FTIR EC (current SPARTAN calibration) vs HIPS Fabs / MAC 10, coloured by season,
          Deming only (OLS dropped at Ann's request, 23/25 Sep), R² + equation per season;
  right = every Addis calibration configuration (Option A) on the Deming slope vs intercept plane,
          with the target (slope 1, intercept 0) and the guardrail passers.

Run with anaconda python from the repo root:
    ~/anaconda3/bin/python deliverables/slide_of_the_week_2026-09-25/build.py

Inputs (the gallery's exported data, same numbers shown in the meeting):
- gallery/app/public/data/filters.json      (unified_filter_dataset.pkl export)
- gallery/app/public/data/calibration.json  (calibration_explorer/cache/batch_results.jsonl export)
"""

from pathlib import Path
import json

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
FIG = HERE / "figures"
FIG.mkdir(exist_ok=True)
DATA = ROOT / "gallery/app/public/data"

GREY, BLUE, PURPLE, ACCENT, AMBER, INK = (
    "#8F8C84", "#2C6E9E", "#7A4FA3", "#B23327", "#C49442", "#22252A")
LIGHT = "#CFCCC4"
DPI = 168
plt.rcParams.update({
    "figure.facecolor": "white", "axes.facecolor": "white",
    "axes.spines.top": False, "axes.spines.right": False,
    "axes.edgecolor": INK, "axes.labelcolor": INK, "xtick.color": INK,
    "ytick.color": INK, "text.color": INK, "font.size": 11,
    "axes.titlesize": 12, "legend.frameon": False,
})
MINUS = "−"


def save(fig, name):
    p = FIG / name
    fig.savefig(p, dpi=DPI, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    Image.open(p).convert("RGB").save(p)
    return p


def deming(x, y, lam):
    mx, my = x.mean(), y.mean()
    sxx, syy = ((x - mx) ** 2).mean(), ((y - my) ** 2).mean()
    sxy = ((x - mx) * (y - my)).mean()
    d = syy - lam * sxx
    m = (d + np.sqrt(d * d + 4 * lam * sxy * sxy)) / (2 * sxy)
    return m, my - m * mx, np.corrcoef(x, y)[0, 1] ** 2


def sgn(v, nd=2):
    return f"{MINUS} {abs(v):.{nd}f}" if v < 0 else f"+ {v:.{nd}f}"


# ---------------------------------------------------------------- data
calib = json.loads((DATA / "calibration.json").read_text())
LAM = calib["deming_lambda_mac10"]          # 2.96, HIPS uncertainty vs IMPROVE CV RMSE
BOX = calib["slope_box"]                    # [0.85, 1.18]
FLOOR = calib["heldout_floor"]              # 0.85 IMPROVE cross-validation R²

rows = pd.DataFrame(json.loads((DATA / "filters.json").read_text())["rows"])
ok = rows[~rows.excluded].dropna(subset=["EC (FTIR)", "HIPS BC"])
addis = ok[ok.site == "Addis Ababa"]
X, Y = "HIPS BC", "EC (FTIR)"

SEASONS = [("Dry (Oct-Feb)", "Dry, Oct to Feb", AMBER),
           ("Belg (Mar-May)", "Belg, Mar to May", PURPLE),
           ("Kiremt (Jun-Sep)", "Kiremt, Jun to Sep", BLUE)]
fits = {"all": deming(addis[X].values, addis[Y].values, LAM)}
for key, _, _ in SEASONS:
    g = addis[addis.season == key]
    fits[key] = deming(g[X].values, g[Y].values, LAM)
means = addis.groupby("season")[[X, Y]].mean()
# what the gallery's per-site panel showed in the meeting (lambda = 1)
site_l1 = {s: deming(g[X].values, g[Y].values, 1.0) for s, g in ok.groupby("site")}

grid = pd.DataFrame(calib["grid"])
G = grid[(grid.tg == "addis") & (grid.mo == "site_heldout") & (grid.lot == "all")
         & (grid.el == "all")].dropna(subset=["dm", "db"])
inbox = G[G.dm.between(*BOX)]
passing = inbox[inbox.ho >= FLOOR]
best_pass = passing.sort_values("db", ascending=False).iloc[0]
near0 = inbox.loc[inbox.db.abs().idxmin()]

# locked anchors (Option A, Lowest-OC/EC 800) must reproduce before anything is shown
a_raw = G[(G.co == "ocec") & (G.cut == 800) & (G.sel == "raw") & (G.sp == "raw")].iloc[0]
a_air = G[(G.co == "ocec") & (G.cut == 800) & (G.sel == "raw") & (G.sp == "airspec")].iloc[0]
assert a_raw.k == 6 and np.allclose([a_raw.om, a_raw.ob, a_raw.ho], [1.585, -3.221, 0.911], atol=1e-3), a_raw
assert a_air.k == 5 and np.allclose([a_air.om, a_air.ob, a_air.dm, a_air.db], [0.857, -1.615, 0.954, -2.089], atol=1e-3), a_air
assert round(site_l1["Addis Ababa"][1], 2) == -6.52 and round(site_l1["JPL"][1], 3) == 0.096

# ---------------------------------------------------------------- figure
fig, (axL, axR) = plt.subplots(1, 2, figsize=(13.4, 5.9), gridspec_kw={"wspace": 0.28})

# left: Addis crossplot by season
hi = float(np.ceil(max(addis[X].max(), addis[Y].max()) + 1.0))
lo_y = min(0.0, fits["all"][1] - 1.2)
xx = np.array([0, hi])
axL.plot(xx, xx, ls=":", color=GREY, lw=1.2, zorder=1)
axL.axhline(0, color=GREY, lw=0.8, zorder=1)
for key, lab, col in SEASONS:
    g = addis[addis.season == key]
    axL.scatter(g[X], g[Y], s=22, color=col, alpha=0.75, lw=0, zorder=3)
    m, b, _ = fits[key]
    axL.plot(xx, m * xx + b, color=col, lw=1.4, zorder=2)
m, b, r2 = fits["all"]
axL.plot(xx, m * xx + b, color=INK, lw=2.0, ls="--", zorder=4)
axL.scatter([0], [b], marker="D", s=46, color=INK, zorder=5, clip_on=False)
axL.set_xlim(0, hi)
axL.set_ylim(lo_y, hi)
axL.set_xlabel("HIPS Fabs / MAC 10 (µg/m³)")
axL.set_ylabel("FTIR EC, current SPARTAN calibration (µg/m³)")
txt = [f"Deming (λ* {LAM:.2f}), n {len(addis)}",
       f"All seasons   y = {m:.2f}x {sgn(b)}   R² {r2:.2f}"]
for key, lab, col in SEASONS:
    mm, bb, rr = fits[key]
    txt.append(f"{lab.split(',')[0]:<12}  y = {mm:.2f}x {sgn(bb)}   R² {rr:.2f}   n {int((addis.season == key).sum())}")
axL.text(0.98, 0.03, "\n".join(txt), transform=axL.transAxes, va="bottom", ha="right",
         fontsize=9.6, family="DejaVu Sans Mono", linespacing=1.45, multialignment="left",
         bbox=dict(boxstyle="round,pad=0.45", fc="white", ec=LIGHT))
cols = {k: c for k, _, c in SEASONS}
for key, lab, col in SEASONS:
    axL.scatter([], [], s=28, color=col, label=lab)
axL.plot([], [], color=INK, ls="--", lw=2, label="All seasons")
axL.plot([], [], color=GREY, ls=":", lw=1.2, label="1:1")
axL.legend(loc="upper left", fontsize=9.5, handletextpad=0.4)
axL.set_title("Addis Ababa (ETAD), Dec 2022 to Sep 2024", loc="left", fontsize=11.5)

# right: slope vs intercept, every configuration
XR, YR = (0, 3.0), (-12, 1.5)
off = int(((G.dm < XR[0]) | (G.dm > XR[1]) | (G.db < YR[0]) | (G.db > YR[1])).sum())
axR.axvspan(*BOX, color=BLUE, alpha=0.07, lw=0)
axR.axhline(0, color=GREY, lw=0.9)
axR.axvline(1, color=GREY, lw=0.9, ls="--")
axR.scatter(G.dm.clip(*XR), G.db.clip(*YR), s=9, color=LIGHT, lw=0, zorder=2,
            label=f"All {len(G):,} configurations")
axR.scatter(passing.dm, passing.db, s=26, color=INK, lw=0, zorder=4,
            label=f"Pass guardrails ({len(passing)})")
axR.scatter([near0.dm], [near0.db], s=90, facecolor="none", edgecolor=ACCENT, lw=1.8, zorder=5)
axR.annotate(f"Closest to 0 in the slope box:\n{near0.dm:.2f}x {sgn(near0.db)}, test-set R² {near0.r2:.2f}",
             (near0.dm, near0.db), xytext=(1.45, 0.45), fontsize=9.5, color=ACCENT,
             arrowprops=dict(arrowstyle="-", color=ACCENT, lw=0.9))
axR.annotate(f"Best passing: {best_pass.dm:.2f}x {sgn(best_pass.db)}\n(Lowest-OC/EC, spline baseline)",
             (best_pass.dm, best_pass.db), xytext=(1.45, -3.6), fontsize=9.5, color=INK,
             arrowprops=dict(arrowstyle="-", color=INK, lw=0.9))
fm, fb, _ = fits["all"]
axR.scatter([fm], [fb], marker="D", s=56, color=INK, zorder=6)
axR.annotate("Current SPARTAN calibration\n(left panel, dashed line)", (fm, fb),
             xytext=(0.95, -8.2), fontsize=9.5,
             arrowprops=dict(arrowstyle="-", color=INK, lw=0.9))
axR.scatter([1], [0], marker="*", s=260, color=ACCENT, edgecolor="white", lw=0.8, zorder=7,
            label="Goal: slope 1, intercept 0")
axR.text((BOX[0] + BOX[1]) / 2, YR[1] - 0.1, f"slope box\n{BOX[0]} to {BOX[1]}", fontsize=8.8, color=BLUE, ha="center", va="top")
axR.set_xlim(*XR)
axR.set_ylim(*YR)
axR.set_xlabel("Addis test set Deming slope")
axR.set_ylabel("Addis test set Deming intercept (µg/m³)")
axR.legend(loc="lower left", fontsize=9.5, handletextpad=0.3, frameon=True, framealpha=1, facecolor="white", edgecolor="white")
axR.set_title("All-season calibrations, Option A, MAC 10", loc="left", fontsize=11.5)
fig_path = save(fig, "f1_addis_intercept_season_and_calibration_plane.png")

# ---------------------------------------------------------------- pptx
INK_RGB = RGBColor(0x22, 0x25, 0x2A)
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
s = prs.slides.add_slide(prs.slide_layouts[6])
tb = s.shapes.add_textbox(Inches(0.45), Inches(0.22), Inches(12.5), Inches(0.95))
tb.text_frame.word_wrap = True
p = tb.text_frame.paragraphs[0]
p.text = ("Addis FTIR EC reads high in Kiremt and low in the dry season, "
          "and no all-season calibration reaches slope 1, intercept 0")
p.font.size, p.font.bold, p.font.color.rgb = Pt(20), True, INK_RGB
im = Image.open(fig_path)
scale = min(12.4 / (im.width / DPI), 5.95 / (im.height / DPI), 1.35)
w = im.width / DPI * scale
s.shapes.add_picture(str(fig_path), Inches((13.333 - w) / 2),
                     Inches(1.25 + (5.95 - im.height / DPI * scale) / 2), Inches(w))

kir, dry = means.loc["Kiremt (Jun-Sep)"], means.loc["Dry (Oct-Feb)"]
say = (
    "This is the week in one slide, and it is the poster story, Addis only. "
    f"On the left is Addis FTIR EC from the current SPARTAN calibration against HIPS absorption divided by MAC 10, "
    f"coloured by season, Deming only. The all-season line has a slope of {fits['all'][0]:.2f} and an intercept of "
    f"about {MINUS}{abs(fits['all'][1]):.0f} micrograms; if FTIR and HIPS were measuring the same carbon that intercept "
    "should be near zero, and at Beijing, Delhi and JPL it is. Follow the colours: Kiremt, the wet season, has the "
    f"highest loadings and FTIR sits above HIPS on average ({kir[Y]:.1f} versus {kir[X]:.1f}); the dry season has lower "
    f"loadings and FTIR sits below ({dry[Y]:.1f} versus {dry[X]:.1f}). So the bias is not one constant offset, it moves "
    "with the season. On the right is every calibration we have fitted for Addis, about nineteen hundred of them, "
    "placed by test set Deming slope and intercept. The red star is where we want to be. The black diamond is the "
    f"current network calibration. Only {len(passing)} configurations pass our guardrails, and the best of them still "
    f"has an intercept of {MINUS}{abs(best_pass.db):.1f}. The one in-box point that gets near zero has a test set R² of "
    f"{near0.r2:.2f}, so it is flat, not right. Takeaway: lumping all seasons together has not got us there, so the "
    "next step is a calibration per season, placed on this same plane."
)
notes = (
    f"Source: gallery exports (filters.json from unified_filter_dataset.pkl; calibration.json from "
    f"calibration_explorer/cache/batch_results.jsonl, exported {calib['generated'][:10]}). Built by build.py. "
    f"Left: {len(addis)} Addis filters with both FTIR EC and HIPS, {addis.date.min()} to {addis.date.max()}, "
    f"excluded filters dropped. Deming λ* {LAM} (HIPS uncertainty vs IMPROVE cross-validation RMSE, MAC 10); "
    f"OLS is not shown, per Ann (25 Sep: data uncertain on both axes, show R² and the Deming equation). "
    f"The gallery panel shown in the meeting uses λ = 1, which gives Addis {site_l1['Addis Ababa'][0]:.2f}x "
    f"{sgn(site_l1['Addis Ababa'][1])}; same λ = 1 per site: Beijing {sgn(site_l1['Beijing'][1])}, Delhi "
    f"{sgn(site_l1['Delhi'][1])}, JPL {sgn(site_l1['JPL'][1], 3)}. Other sites are deliberately off the slide (poster "
    "is Addis only). Seasonal intercepts are each less negative than the all-season one because each season's line "
    "is fitted over a narrower range; the point is that the seasons sit on different lines. "
    f"Right: Option A (site-grouped 5-fold CV, first major minimum, site-disjoint 80/20 fit), Addis test set, "
    f"fixed 190-filter set, MAC 10, {len(G):,} configurations across cohorts (Lowest-OC/EC, spectral analogs, "
    f"Ethiopia-shaped smoke, biomass smoke, entire IMPROVE network) and spectra (raw, spline baseline, second "
    f"derivative). {off} extreme configurations are pinned to the plot edge. Guardrails = Deming slope in "
    f"{BOX[0]} to {BOX[1]} and IMPROVE cross-validation R² at least {FLOOR}. Best passing: {best_pass.co} "
    f"{int(best_pass.cut)}, {best_pass.sp}, k {int(best_pass.k)}, {best_pass.dm:.3f}x {sgn(best_pass.db, 3)}, test set "
    f"R² {best_pass.r2:.2f}. Closest to zero in the slope box: {near0.co} {int(near0.cut)}, {near0.sp}, k "
    f"{int(near0.k)}, {near0.dm:.3f}x {sgn(near0.db, 3)}, test set R² {near0.r2:.3f}, IMPROVE cross-validation R² "
    f"{near0.ho:.3f}. Anchors reproduced before plotting: Lowest-OC/EC 800 raw k 6, OLS 1.585x {MINUS} 3.221, "
    "held-out TOR R² 0.911; with spline baseline (AIRSpec) k 5, Deming 0.95x {MINUS} 2.09. "
    "Caveats: analog cutoffs resolve to fewer filters than nominal (top 500 gives 477) because the cutoff is applied "
    "before TOR eligibility; second derivative configurations are still in the cloud and will be dropped; the VIBES "
    "comparison shown in the meeting cut at 1500 cm⁻¹ and is being rerun over the full range, so VIBES points are not "
    "on this slide. Agreed wording from this meeting: 'test set' (not held-out TOR), 'cross-validation protocol', "
    "'spline baseline' and 'raw spectra'. Blocked: IMPROVE data download from the FED site has failed for four days; "
    "support emailed, Ann can chase a contact if it stays down."
).replace("{MINUS}", MINUS)
s.notes_slide.notes_text_frame.text = f"SAY: {say}\n\nNOTES: {notes}"

out = HERE / "slide_of_the_week_2026-09-25.pptx"
prs.save(out)

# ---------------------------------------------------------------- checks
chk = Presentation(out)
assert len(chk.slides) == 1
for sl in chk.slides:
    txt = " ".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame)
    nt = sl.notes_slide.notes_text_frame.text
    assert chr(0x2014) not in txt + nt, "em dash on slide"
    assert nt.startswith("SAY:") and "NOTES:" in nt, "notes missing"
    assert sum(1 for sh in sl.shapes if sh.shape_type == 13) == 1, "figure missing"
assert chr(0x2014) not in Path(__file__).read_text(), "em dash in builder"
print(f"saved {out.relative_to(ROOT)}: 1 slide, 1 figure, notes present, no em dashes")
print(f"all-season Deming {fits['all'][0]:.3f}x {fits['all'][1]:+.3f}; passing {len(passing)}; "
      f"best passing {best_pass.db:.3f}; near0 r2 {near0.r2:.3f}; off-canvas {off}")
