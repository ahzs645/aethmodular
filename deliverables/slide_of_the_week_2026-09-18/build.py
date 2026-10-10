"""Slide of the week, 18 Sep 2026 (Ann + Satoshi call): figure + pptx, words over numbers.

One slide, one figure:
  left  = every Addis calibration configuration (Option A) placed by its Deming intercept, one row
          per spectral preprocessing (raw, second derivative, spline baseline), with the entire-network
          calibration marked in each row: the baseline is the lever that moves the intercept;
  right = Satoshi's test for next week, as three labelled cards (no numbers).

Run with anaconda python from the repo root:
    ~/anaconda3/bin/python deliverables/slide_of_the_week_2026-09-18/build.py

Input: gallery/app/public/data/calibration.json (calibration_explorer/cache/batch_results.jsonl export).
"""

from pathlib import Path
import json

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
import numpy as np
import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
FIG = HERE / "figures"
FIG.mkdir(exist_ok=True)
DATA = ROOT / "gallery/app/public/data"

GREY, BLUE, PURPLE, ACCENT, AMBER, INK = (
    "#8F8C84", "#2C6E9E", "#7A4FA3", "#B23327", "#C49442", "#22252A")
DPI = 168
plt.rcParams.update({
    "figure.facecolor": "white", "axes.facecolor": "white",
    "axes.spines.top": False, "axes.spines.right": False,
    "axes.edgecolor": INK, "axes.labelcolor": INK, "xtick.color": INK,
    "ytick.color": INK, "text.color": INK, "font.size": 11,
    "axes.titlesize": 12.5, "legend.frameon": False,
})
MINUS = "−"


def save(fig, name):
    p = FIG / name
    fig.savefig(p, dpi=DPI, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    Image.open(p).convert("RGB").save(p)
    return p


def sgn(v, nd=2):
    return f"{MINUS}{abs(v):.{nd}f}" if v < 0 else f"{v:.{nd}f}"


# ---------------------------------------------------------------- data
calib = json.loads((DATA / "calibration.json").read_text())
g = pd.DataFrame(calib["grid"])
G = g[(g.tg == "addis") & (g.mo == "site_heldout") & (g.lot == "all") & (g.el == "all")].dropna(subset=["db"])
pool = G[G.co == "pool"].set_index("sp")
# the meeting's headline: entire-network calibration, raw vs spline baseline (Option A, MAC 10)
assert np.allclose([pool.loc["raw", "db"], pool.loc["airspec", "db"]], [-4.8196, -0.7836], atol=1e-4)
ROWS = [("raw", "Raw spectra", GREY, "Far from zero"),
        ("deriv2", "Second derivative", PURPLE, "In between, widely spread"),
        ("airspec", "Spline baseline", BLUE, "Closest to zero")]
stats = {sp: G[G.sp == sp].db for sp, *_ in ROWS}
near = {sp: float((s.abs() < 1).mean()) for sp, s in stats.items()}
med = {sp: float(s.median()) for sp, s in stats.items()}
assert med["airspec"] > med["deriv2"] > med["raw"]

# ---------------------------------------------------------------- figure
fig = plt.figure(figsize=(14.6, 6.0))
gs = fig.add_gridspec(1, 2, width_ratios=[1.75, 1], wspace=0.06)
ax = fig.add_subplot(gs[0])
card = fig.add_subplot(gs[1])

XL = (-12, 2)
rng = np.random.default_rng(20260918)
ax.axvline(0, color=ACCENT, lw=1.4, zorder=1)
ax.text(0.15, 2.62, "goal:\nintercept 0", color=ACCENT, fontsize=9.8, va="top")
off = 0
for i, (sp, name, col, word) in enumerate(ROWS):
    s = stats[sp].to_numpy()
    off += int(((s < XL[0]) | (s > XL[1])).sum())
    yy = i + rng.uniform(-0.28, 0.28, len(s))
    ax.scatter(np.clip(s, *XL), yy, s=9, color=col, alpha=0.45, lw=0, zorder=2)
    ax.text(XL[0] - 0.3, i + 0.08, name, ha="right", va="center", fontsize=11.5, color=col, fontweight="bold")
    ax.text(XL[0] - 0.3, i - 0.17, word, ha="right", va="center", fontsize=9.8, color=INK)
    ax.scatter([pool.loc[sp, "db"]], [i], marker="D", s=70, color=INK, zorder=4)
ax.annotate("", (pool.loc["airspec", "db"], 2 - 0.12), (pool.loc["raw", "db"], 0 + 0.12),
            arrowprops=dict(arrowstyle="-|>", color=INK, lw=1.4, connectionstyle="arc3,rad=-0.25"))
ax.text(-11.9, 1.5, "Same entire-network calibration: the baseline alone moves it\n"
        "most of the way to zero, but its slope drops well below 1:1", fontsize=10, ha="left",
        va="center", color=INK)
ax.scatter([], [], marker="D", s=50, color=INK, label="Entire IMPROVE network calibration")
ax.scatter([], [], s=14, color=GREY, label="Every other configuration tried")
ax.legend(loc="lower left", fontsize=9.6, bbox_to_anchor=(0.0, -0.2), ncol=2)
ax.set_xlim(*XL)
ax.set_ylim(-0.6, 2.75)
ax.set_yticks([])
ax.spines["left"].set_visible(False)
ax.set_xlabel("Addis Deming intercept, test set (µg/m³)  ← more negative")
ax.set_title("Where every Addis calibration lands, by preprocessing", loc="left", fontsize=12)
for i, (sp, *_rest) in enumerate(ROWS):
    tag, colr = ("dropped", GREY) if sp != "airspec" else ("kept", BLUE)
    ax.text(XL[1] - 0.1, i, tag, ha="right", va="center", fontsize=10.5, color=colr, fontstyle="italic",
            bbox=dict(boxstyle="round,pad=0.25", fc="white", ec=colr, lw=1))

# right: Satoshi's test
card.set_xlim(0, 1)
card.set_ylim(0, 1)
card.axis("off")
card.set_title("Next: do the IMPROVE analogs show the same bias?", loc="left", fontsize=12)
items = [
    ("TOR EC vs HIPS", "is the reference itself off\nfor these kinds of filters?"),
    ("FTIR EC vs TOR EC", "does the model already miss\non held-out IMPROVE analogs?"),
    ("FTIR EC vs HIPS", "the same picture as Addis?"),
]
ys = [0.79, 0.555, 0.32]
for (head, body), yy in zip(items, ys):
    card.add_patch(FancyBboxPatch((0.03, yy - 0.085), 0.94, 0.185,
                                  boxstyle="round,pad=0.01,rounding_size=0.03", fc="white", ec=BLUE, lw=1.4))
    # tiny cross-plot icon
    ix, iy, w = 0.08, yy - 0.06, 0.12
    card.plot([ix, ix, ix + w], [iy + w, iy, iy], color=INK, lw=1)
    card.plot([ix, ix + w], [iy, iy + w], color=GREY, lw=0.9, ls=":")
    card.plot([ix + 0.01, ix + w], [iy + 0.005, iy + w * 0.55], color=BLUE, lw=1.4)
    card.text(0.25, yy + 0.04, head, fontsize=11, fontweight="bold", color=BLUE, va="center")
    card.text(0.25, yy - 0.035, body, fontsize=10, color=INK, va="center")
card.text(0.03, 0.17, "For each season's analog set, and coloured by the\nTOR pyrolysis (OP) fraction",
          fontsize=10, color=INK, va="center")
card.text(0.03, 0.04, "Also settled: CO₂ band and the top end out of the\nmatch; Satoshi's group to baseline "
          "our Addis spectra", fontsize=10, color=GREY, va="center")
fig_path = save(fig, "f1_baseline_is_the_lever_and_next_test.png")

# ---------------------------------------------------------------- pptx
INK_RGB = RGBColor(0x22, 0x25, 0x2A)
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
sl = prs.slides.add_slide(prs.slide_layouts[6])
tb = sl.shapes.add_textbox(Inches(0.45), Inches(0.22), Inches(12.5), Inches(0.95))
tb.text_frame.word_wrap = True
p = tb.text_frame.paragraphs[0]
p.text = ("Baseline correction is the lever that moves the Addis intercept; next, check whether "
          "the IMPROVE analogs carry the same bias")
p.font.size, p.font.bold, p.font.color.rgb = Pt(20), True, INK_RGB
im = Image.open(fig_path)
scale = min(12.4 / (im.width / DPI), 6.0 / (im.height / DPI), 1.35)
w = im.width / DPI * scale
sl.shapes.add_picture(str(fig_path), Inches((13.333 - w) / 2),
                      Inches(1.25 + (6.0 - im.height / DPI * scale) / 2), Inches(w))

pr, pa, pd2 = pool.loc["raw"], pool.loc["airspec"], pool.loc["deriv2"]
say = (
    "This is where the call with Satoshi landed. On the left, every dot is one Addis calibration we have "
    "tried, placed by its intercept, with zero on the red line. The rows split them by how the spectra "
    "were prepared. Raw spectra sit far to the left. Second derivative is in between and spread out. The "
    "spline baseline sits closest to zero. The black diamond is the same entire-network calibration in each "
    "row: changing only the baseline moves it most of the way to zero, but its slope falls well below one. "
    "So we decided to drop raw and second derivative and use the spline baseline from here on, which takes "
    "a big lever out of the search. On the right is Satoshi's question for next week. The analogs we pick "
    "look a lot like Addis, so do they show the same bias? For each season's analog set I will plot TOR EC "
    "against HIPS, FTIR against TOR on held-out analogs, and FTIR against HIPS, coloured by the pyrolysis "
    "fraction. If the analogs miss the same way, the problem is in the reference or the spectra, not in "
    "how we choose filters."
)
notes = (
    "Source: gallery calibration.json (calibration_explorer batch grid), Addis test set, Option A "
    "(site-grouped 5-fold CV, first major minimum, site-disjoint 80/20 fit), MAC 10, Deming. "
    f"{len(stats['raw'])} configurations per row across cohorts and cutoffs; {off} beyond the plotted range "
    f"are pinned to the edge. Median intercept: raw {sgn(med['raw'])}, second derivative {sgn(med['deriv2'])}, "
    f"spline baseline {sgn(med['airspec'])}; share within ±1 µg/m³: {near['raw']:.0%}, {near['deriv2']:.0%}, "
    f"{near['airspec']:.0%}. Entire IMPROVE network calibration: raw {pr.dm:.2f}x {sgn(pr.db)} (k {int(pr.k)}), "
    f"spline baseline {pa.dm:.2f}x {sgn(pa.db)} (k {int(pa.k)}), second derivative {pd2.dm:.2f}x {sgn(pd2.db)}; "
    "these are the numbers quoted in the call. The biomass-smoke calibration export still showed its "
    "baseline incorrectly and is left out. Decisions in the call: spline baseline only (no raw, no second "
    "derivative); exclude the CO2 band and the top end (3500 or 3600 cm-1, nine filters apart, pick one) "
    "from matching; keep exclusions identical when the new baseline arrives. Satoshi's asks: (1) for each "
    "season's analog set, TOR EC vs HIPS Fabs/10, FTIR EC vs TOR EC on held-out analogs, FTIR EC vs Fabs/10, "
    "and where the analogs sit when predicted by the same model as Addis; (2) consider per-filter nearest "
    "analogs instead of matching the season median (less urgent, since spectra are homogeneous within a "
    "season); (3) send raw Addis spectra so Francois Kamper can baseline them with their new method, "
    "runnable code hopefully next week. Ann's asks: get TOR fractions (OPR, OPT, EC1 to EC3) and colour "
    "crossplots by OP; fix how to cross-validate so it stops being a lever. Discussion: Kiremt Addis spectra "
    "sit below most of the training range, which may explain the poor intercept there; extrapolating down "
    "from higher-loaded training filters is usually fine, extrapolating up is not. Addis is unique in the "
    "intercept; Bujumbura is unusual only in slope. Ann and Ahmad's afternoon meeting moved to next week."
)
sl.notes_slide.notes_text_frame.text = f"SAY: {say}\n\nNOTES: {notes}"

out = HERE / "slide_of_the_week_2026-09-18.pptx"
prs.save(out)

# ---------------------------------------------------------------- checks
chk = Presentation(out)
assert len(chk.slides) == 1
for s_ in chk.slides:
    txt = " ".join(sh.text_frame.text for sh in s_.shapes if sh.has_text_frame)
    nt = s_.notes_slide.notes_text_frame.text
    assert chr(0x2014) not in txt + nt, "em dash on slide"
    assert nt.startswith("SAY:") and "NOTES:" in nt, "notes missing"
    assert sum(1 for sh in s_.shapes if sh.shape_type == 13) == 1, "figure missing"
assert chr(0x2014) not in Path(__file__).read_text(), "em dash in builder"
print(f"saved {out.relative_to(ROOT)}: 1 slide, 1 figure, notes present, no em dashes")
print("medians", med, "near", near, "off", off)
