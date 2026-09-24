"""Stream both IMPROVE portal workbooks into complete, queryable Parquet tables."""

from __future__ import annotations

import argparse
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import re

from openpyxl import load_workbook
import pyarrow as pa
import pyarrow.parquet as pq


HERE = Path(__file__).resolve().parent
REPO = HERE.parents[3]
MANIFEST = HERE / "improve_raw_parquet.manifest.json"
CHUNK_ROWS = 20_000


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def filename(sheet: str, kind: str | None = None) -> str:
    stem = re.sub(r"[^a-z0-9]+", "_", sheet.lower()).strip("_")
    return f"improve_{kind + '_' if kind else ''}{stem}.parquet"


def worksheet_rows(sheet):
    iterator = sheet.iter_rows(values_only=True)
    headers = next(iterator)
    if not headers or any(not isinstance(header, str) or not header for header in headers):
        raise ValueError(f"Invalid headers in {sheet.title}")
    return list(headers), iterator


def convert_data(sheet, target: Path, expected_rows: int) -> dict:
    headers, iterator = worksheet_rows(sheet)
    if headers[:5] != ["Dataset", "SiteCode", "POC", "Date", "AuxID"] or not all(
        name.endswith("_Val") for name in headers[5:]
    ):
        raise ValueError(f"Unexpected IMPROVE Data columns: {headers}")
    fields = [
        pa.field("Dataset", pa.string()), pa.field("SiteCode", pa.string()),
        pa.field("POC", pa.int64()), pa.field("Date", pa.string()),
        pa.field("AuxID", pa.int64()),
        *[pa.field(name, pa.float64()) for name in headers[5:]],
        pa.field("sample_date", pa.date32()),
    ]
    schema = pa.schema(fields)
    total = 0
    batch = []
    with pq.ParquetWriter(target, schema, compression="zstd", use_dictionary=True) as writer:
        for row in iterator:
            if len(row) != len(headers):
                raise ValueError(f"Unexpected Data row width at {total + 2}")
            dataset, site, poc, original_date, aux, *numbers = row
            if not isinstance(original_date, str):
                raise ValueError(f"Unexpected Date value at row {total + 2}")
            parsed_date = datetime.strptime(original_date, "%m/%d/%Y").date()
            converted = [dataset, site, poc, original_date, aux]
            for value in numbers:
                if value in (None, ""):
                    converted.append(None)
                elif isinstance(value, (float, int)):
                    converted.append(float(value))
                else:
                    raise ValueError(f"Unexpected numeric value at row {total + 2}: {value!r}")
            batch.append([*converted, parsed_date])
            total += 1
            if len(batch) == CHUNK_ROWS:
                writer.write_table(pa.Table.from_pylist([dict(zip(schema.names, values)) for values in batch], schema=schema))
                batch.clear()
        if batch:
            writer.write_table(pa.Table.from_pylist([dict(zip(schema.names, values)) for values in batch], schema=schema))
    if total != expected_rows or pq.ParquetFile(target).metadata.num_rows != expected_rows:
        raise ValueError(f"Data row count mismatch: {total} vs {expected_rows}")
    return {"fileName": target.name, "rows": total, "columns": schema.names, "sha256": sha256(target), "bytes": target.stat().st_size}


def metadata_value(value):
    if value in (None, ""):
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def convert_metadata(sheet, target: Path) -> dict:
    headers, iterator = worksheet_rows(sheet)
    rows = [tuple(metadata_value(value) for value in row) for row in iterator]
    columns = list(zip(*rows)) if rows else [[] for _ in headers]
    fields = []
    for name, values in zip(headers, columns):
        present = [value for value in values if value is not None]
        if present and all(isinstance(value, int) and not isinstance(value, bool) for value in present):
            dtype = pa.int64()
        elif present and all(isinstance(value, (int, float)) and not isinstance(value, bool) for value in present):
            dtype = pa.float64()
        else:
            dtype = pa.string()
        fields.append(pa.field(name, dtype))
    schema = pa.schema(fields)
    normalized = []
    for row in rows:
        normalized.append({name: (str(value) if value is not None and pa.types.is_string(schema.field(name).type) else value)
                           for name, value in zip(headers, row)})
    pq.write_table(pa.Table.from_pylist(normalized, schema=schema), target, compression="zstd")
    if pq.ParquetFile(target).metadata.num_rows != len(rows):
        raise ValueError(f"Metadata row count mismatch in {sheet.title}")
    return {"fileName": target.name, "rows": len(rows), "columns": headers, "sha256": sha256(target), "bytes": target.stat().st_size}


def sheets_equal(left, right) -> bool:
    if left.max_row != right.max_row or left.max_column != right.max_column:
        return False
    return all(a == b for a, b in zip(left.iter_rows(values_only=True), right.iter_rows(values_only=True)))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path)
    args = parser.parse_args()
    manifest = json.loads(MANIFEST.read_text())
    if (manifest.get("schemaVersion"), manifest.get("id"), manifest.get("script")) != (1, "aeth.improve-raw-parquet", Path(__file__).name):
        raise ValueError("Unexpected recipe manifest")
    output_dir = args.output_dir or REPO / manifest["outputFromRepo"]
    output_dir.mkdir(parents=True, exist_ok=True)
    sources = {}
    books = {}
    for kind, spec in manifest["sources"].items():
        source = args.source_dir / spec["fileName"]
        digest = sha256(source)
        if digest != spec["sha256"]:
            raise ValueError(f"Source checksum changed: {source.name}")
        sources[kind] = {"fileName": source.name, "sha256": digest, "bytes": source.stat().st_size}
        books[kind] = load_workbook(source, read_only=True, data_only=True)
    chem, laser = books["chem"], books["laser"]
    if set(chem.sheetnames) != set(laser.sheetnames) or set(chem.sheetnames) != {"Data", *manifest["sharedSheets"], *manifest["distinctSheets"]}:
        raise ValueError("Unexpected IMPROVE workbook sheets")
    for sheet in manifest["sharedSheets"]:
        if not sheets_equal(chem[sheet], laser[sheet]):
            raise ValueError(f"Shared metadata differs between workbooks: {sheet}")
    outputs = []
    for kind, book in books.items():
        result = convert_data(book["Data"], output_dir / filename("Data", kind), manifest["sources"][kind]["dataRows"])
        outputs.append({"source": kind, "sheet": "Data", **result})
    for sheet in manifest["sharedSheets"]:
        outputs.append({"source": "both", "sheet": sheet, **convert_metadata(chem[sheet], output_dir / filename(sheet))})
    for kind, book in books.items():
        for sheet in manifest["distinctSheets"]:
            outputs.append({"source": kind, "sheet": sheet, **convert_metadata(book[sheet], output_dir / filename(sheet, kind))})
    for book in books.values():
        book.close()
    receipt = {
        "recipeId": manifest["id"], "recipeVersion": manifest["version"],
        "manifestSha256": sha256(MANIFEST), "scriptSha256": sha256(Path(__file__)),
        "sources": sources, "outputs": outputs,
        "sharedSheetsVerifiedEqual": manifest["sharedSheets"],
        "valueRules": ["Data numeric blanks become null", "Date source text is retained", "sample_date is a parsed date", "Shared metadata sheets are stored once"],
    }
    (output_dir / "improve_raw.provenance.json").write_text(json.dumps(receipt, indent=2) + "\n")
    print(json.dumps({"tables": len(outputs), "rows": sum(item["rows"] for item in outputs), "bytes": sum(item["bytes"] for item in outputs), "output": str(output_dir)}))


if __name__ == "__main__":
    main()
