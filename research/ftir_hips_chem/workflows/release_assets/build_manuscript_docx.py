"""Render the scientific manuscript JSON as a restrained Word document."""

from pathlib import Path
import json
import sys

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(sys.argv[1]).resolve()
content = json.loads((root / "manuscript_content.json").read_text())
doc = Document()
# Remove decorative borders inherited from the installed default template.
for tree in [doc.styles.element, doc.element]:
    for border in list(tree.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
sec = doc.sections[0]
sec.top_margin = Inches(0.7)
sec.bottom_margin = Inches(0.7)
sec.left_margin = Inches(0.75)
sec.right_margin = Inches(0.75)
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption"]:
    st = doc.styles[name]
    st.font.name = "Arial"
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_after = Pt(7)
    if name == "Normal":
        st.font.size = Pt(10.5)
        st.paragraph_format.line_spacing = 1.12
    if name == "Title":
        st.font.size = Pt(21)
        st.font.bold = True
        st.paragraph_format.space_after = Pt(10)
    if name == "Subtitle":
        st.font.size = Pt(11)
    if name == "Heading 1":
        st.font.size = Pt(14)
        st.paragraph_format.space_before = Pt(12)
    if name == "Heading 2":
        st.font.size = Pt(11.5)
        st.font.bold = True
        st.paragraph_format.space_before = Pt(9)
    if name == "Caption":
        st.font.size = Pt(9)
        st.font.bold = False
        st.paragraph_format.line_spacing = 1.05
# Page numbers aid review of a scientific manuscript.
f = sec.footer.paragraphs[0]
f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
page_run = OxmlElement("w:r")
page_text = OxmlElement("w:t")
page_text.text = "1"
page_run.append(page_text)
field.append(page_run)
f._p.append(field)


def add_table(item):
    p = doc.add_paragraph(item["text"], "Caption")
    p.paragraph_format.keep_with_next = True
    tab = doc.add_table(rows=1, cols=len(item["columns"]))
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.autofit = False
    widths = (
        [1.1, 0.75, 0.7, 0.8, 3.65]
        if len(item["columns"]) == 5
        else [1.0, 1.15, 0.5, 1.0, 1.15, 1.1, 1.1]
        if len(item["columns"]) == 7
        else [1.35, 1.8, 1.925, 1.925]
    )
    scale = 7 / sum(widths)
    widths = [x * scale for x in widths]
    for j, width in enumerate(widths):
        tab.columns[j].width = Inches(width)
    for j, label in enumerate(item["columns"]):
        tab.rows[0].cells[j].text = label
    for data in item["rows"]:
        row = tab.add_row()
        for j, value in enumerate(data):
            row.cells[j].text = str(value)
    for ri, row in enumerate(tab.rows):
        trPr = row._tr.get_or_add_trPr()
        if ri == 0:
            repeat = OxmlElement("w:tblHeader")
            trPr.append(repeat)
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)
        for j, cell in enumerate(row.cells):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for side in ["top", "left", "bottom", "right"]:
                el = OxmlElement("w:" + side)
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), "4")
                el.set(qn("w:color"), "D9D9D9")
                borders.append(el)
            props.append(borders)
            margins = OxmlElement("w:tcMar")
            for side in ["top", "left", "bottom", "right"]:
                el = OxmlElement("w:" + side)
                el.set(qn("w:w"), "80")
                el.set(qn("w:type"), "dxa")
                margins.append(el)
            props.append(margins)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "E8EDF2" if ri == 0 else "FFFFFF")
            props.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.line_spacing = 1
                # Each summary table fits on one page; keep its sites together.
                p.paragraph_format.keep_with_next = ri < len(tab.rows) - 1
                p.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT
                    if j < 2 or len(widths) == 5 and j == 4
                    else WD_ALIGN_PARAGRAPH.CENTER
                )
                for r in p.runs:
                    r.font.size = Pt(8.3)
                    r.bold = ri == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


for item in content:
    kind = item["kind"]
    if kind == "title":
        doc.add_paragraph(item["text"], "Title")
    elif kind == "subtitle":
        doc.add_paragraph(item["text"], "Subtitle")
    elif kind == "heading":
        p = doc.add_paragraph(item["text"], "Heading 1")
        if item["text"] == "Results":
            p.paragraph_format.page_break_before = True
    elif kind == "subheading":
        doc.add_paragraph(item["text"], "Heading 2")
    elif kind == "table":
        add_table(item)
    elif kind == "figure":
        doc.add_page_break()
        doc.add_paragraph(item["text"], "Heading 1")
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        run = p.add_run()
        run.add_picture(str(root / item["path"]), width=Inches(7))
        # Accessible image description retains the figure's scientific role.
        for drawing in run._r.xpath(".//wp:docPr"):
            drawing.set("descr", item["caption"])
        doc.add_paragraph(item["caption"], "Caption")
    elif kind == "reference":
        p = doc.add_paragraph(item["text"])
        p.paragraph_format.line_spacing = 1
        for r in p.runs:
            r.font.size = Pt(8.5)
    else:
        doc.add_paragraph(item["text"])
doc.core_properties.title = content[0]["text"]
doc.core_properties.subject = "Consolidated filter-only Methods and Results"
doc.core_properties.author = ""
doc.save(root / "scientific_draft.docx")
print(root / "scientific_draft.docx")
