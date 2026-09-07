"""One figure per slide, white, nothing else: the plain figure deck.

Requested 2026-09-01 as the companion to the HTML gallery: every figure from
the session on its own slide, no design, so the graphs can be dropped straight
into any other deck. Titles and captions come from build_figure_gallery.CAPTIONS
so the deck, the gallery and the worklog all say the same thing about each
figure. Figures without an authored caption get their filename.

Repo rule (memory: deck-no-em-dashes): no em dash may reach a .pptx. The build
asserts it rather than trusting the author.

    python3 build_figure_deck.py                       # the four session sets
    python3 build_figure_deck.py --all                 # every figure in the tree
    python3 build_figure_deck.py --sets stability,decisions --out deck.pptx
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from PIL import Image

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from build_figure_gallery import CAPTIONS, PLOTS, collect, title_from  # noqa: E402

# the narrative order for the session sets: the result, what the search can and
# cannot deliver, the findings that survived, then the eliminations
SESSION_SETS = ["crossplots", "decisions", "stability", "pathways"]

SLIDE_W, SLIDE_H = 13.333, 7.5           # 16:9 inches
MARGIN = 0.5
TITLE_H = 0.75
CAPTION_H = 0.7
INK = RGBColor(0x22, 0x25, 0x2A)
MUTED = RGBColor(0x6E, 0x71, 0x78)
FONT = "Calibri"                          # ships with Office, renders true in QA


def no_em_dash(text: str, where: str) -> str:
    if "—" in text:
        raise SystemExit(f"em dash in {where}: {text!r}  (repo rule: none in decks)")
    return text


def add_figure_slide(prs, png: Path, title: str, caption: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])        # blank

    tb = slide.shapes.add_textbox(Inches(MARGIN), Inches(MARGIN),
                                  Inches(SLIDE_W - 2 * MARGIN), Inches(TITLE_H))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    run = tf.paragraphs[0].add_run()
    run.text = no_em_dash(title, png.name)
    run.font.name = FONT
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = INK

    # fit the image into the band between title and caption, preserving aspect
    box_top = MARGIN + TITLE_H + 0.15
    box_h = SLIDE_H - box_top - MARGIN - CAPTION_H - 0.1
    box_w = SLIDE_W - 2 * MARGIN
    with Image.open(png) as im:
        w_px, h_px = im.size
    aspect = w_px / h_px
    if box_w / box_h > aspect:            # band is wider than the image
        h = box_h; w = h * aspect
    else:
        w = box_w; h = w / aspect
    left = (SLIDE_W - w) / 2
    top = box_top + (box_h - h) / 2
    slide.shapes.add_picture(str(png), Inches(left), Inches(top),
                             width=Inches(w), height=Inches(h))

    cb = slide.shapes.add_textbox(Inches(MARGIN), Inches(SLIDE_H - MARGIN - CAPTION_H),
                                  Inches(SLIDE_W - 2 * MARGIN), Inches(CAPTION_H))
    cf = cb.text_frame
    cf.word_wrap = True
    cf.margin_left = cf.margin_right = cf.margin_top = cf.margin_bottom = 0
    crun = cf.paragraphs[0].add_run()
    crun.text = no_em_dash(caption, png.name)
    crun.font.name = FONT
    crun.font.size = Pt(12)
    crun.font.color.rgb = MUTED
    return slide


def build(groups, out: Path):
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    n = 0
    for name, pngs in groups:
        for png in pngs:
            rel = f"{name}/{png.name}"
            cap = CAPTIONS.get(rel)
            title = cap[0] if cap else title_from(png.stem)
            body = cap[1] if cap else f"{rel}  (no authored caption; see the builder script)"
            add_figure_slide(prs, png, title, body)
            n += 1
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print(f"  {n} slides -> {out}  ({out.stat().st_size / 1e6:.1f} MB)")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--sets", help="comma-separated figure directories")
    ap.add_argument("--all", action="store_true")
    ap.add_argument("--out", default=str(HERE.parents[2] /
                    "deliverables/figure_deck_2026-09-01/figure_deck.pptx"))
    a = ap.parse_args()
    if a.all:
        groups = collect(None)
    else:
        sets = [s.strip() for s in a.sets.split(",")] if a.sets else SESSION_SETS
        found = dict(collect(sets))
        groups = [(s, found[s]) for s in sets if s in found]     # keep narrative order
    build(groups, Path(a.out))


if __name__ == "__main__":
    main()
